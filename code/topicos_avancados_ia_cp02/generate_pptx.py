import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()
# Set 16:9 aspect ratio
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# 1. Add Title (Name)
txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(6.0), Inches(1.5))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Cauã da Cunha Lira"
p.font.size = Pt(40)
p.font.name = "Calibri"
p.font.color.rgb = RGBColor(51, 51, 51)

# 2. Add content (Modelo, Dataset, Resultados, Link)
txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.5))
tf2 = txBox2.text_frame
tf2.word_wrap = True

def add_info(bold_text, normal_text):
    p = tf2.add_paragraph()
    p.space_after = Pt(20)
    run1 = p.add_run()
    run1.text = bold_text + " "
    run1.font.bold = True
    run1.font.size = Pt(22)
    run1.font.name = "Calibri"
    
    run2 = p.add_run()
    run2.text = normal_text
    run2.font.bold = False
    run2.font.size = Pt(22)
    run2.font.name = "Calibri"

add_info("Modelo:", "RAPUNet-Zeta (CNN - Encoder ResNet50V2 + Decoder Spatial Attention + Mish)")
add_info("Dataset:", "Kvasir-SEG (Pólipos Gastrointestinais, 1.000 imagens médicas)")
add_info("Resultados:", "accuracy: 0.9627 - loss: 0.3416 - dice coef: 0.6536")

p_link = tf2.add_paragraph()
p_link.space_after = Pt(20)
r_link_b = p_link.add_run()
r_link_b.text = "Link:\n"
r_link_b.font.bold = True
r_link_b.font.size = Pt(22)

r_link = p_link.add_run()
r_link.text = "https://github.com/lira-labs/workspace/tree/main/topicos-avancados-ia-cp02"
r_link.font.size = Pt(20)
r_link.font.color.rgb = RGBColor(0, 86, 179)
r_link.font.underline = True

# 3. Add Images on the right side
# Image 1 (Training curves)
img1_path = "plot_results/training_curves.png"
slide.shapes.add_picture(img1_path, Inches(6.8), Inches(0.5), width=Inches(5.5))

# Image 2 (Segmentation Results)
img2_path = "plot_results/segmentation_results.png"
slide.shapes.add_picture(img2_path, Inches(6.8), Inches(4.0), width=Inches(5.5))

prs.save("slide_apresentacao_cp02.pptx")
print("PPTX gerado com sucesso!")
